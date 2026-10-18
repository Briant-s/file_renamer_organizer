from pathlib import Path

from pptx import Presentation
from pptx.exc import PackageNotFoundError

from tifo.extractors.models import ExtractedFile

# Cap the amount of text we pull out.
MAX_CHARS = 2000


def _iter_shape_text(shape) -> list[str]:
    """Pull text from a single shape: text frames and tables."""
    texts: list[str] = []

    if shape.has_text_frame:
        frame_text = shape.text_frame.text.strip()
        if frame_text:
            texts.append(frame_text)

    if shape.has_table:
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                texts.append(" | ".join(cells))

    return texts


def extract_pptx(file_path: Path) -> ExtractedFile:
    """Extract slide text from a .pptx deck.

    Returns content=None (rather than raising) for corrupt or non-OOXML files
    so a single bad deck never breaks a whole extraction run. Slide titles are
    collected first since they're usually the most descriptive text.
    """
    try:
        prs = Presentation(str(file_path))
    except (PackageNotFoundError, KeyError, ValueError, OSError):
        # Corrupt, encrypted, or not actually a valid OOXML .pptx
        return ExtractedFile(content=None, summary=None, metadata={})

    titles: list[str] = []
    body: list[str] = []

    for slide in prs.slides:
        # Slide title (if the layout defines one) is the strongest signal.
        title_shape = slide.shapes.title
        title_text = title_shape.text.strip() if title_shape is not None else ""
        if title_text:
            titles.append(title_text)
        # python-pptx hands back a fresh wrapper object on each `.title` access,
        # so identity (`is`) never matches — compare the stable shape_id instead.
        title_id = title_shape.shape_id if title_shape is not None else None

        for shape in slide.shapes:
            # Skip the title shape here; already captured above.
            if title_id is not None and shape.shape_id == title_id:
                continue
            body.extend(_iter_shape_text(shape))

    # Titles first, then body text, so the most descriptive content leads.
    full_text = "\n".join(titles + body).strip()

    return ExtractedFile(
        content=full_text[:MAX_CHARS] if full_text else None,
        summary=None,
        metadata={"slide_count": len(prs.slides)},
    )
