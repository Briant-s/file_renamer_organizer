import openpyxl
from pathlib import Path
from src.extractors.models import ExtractedFile

MAX_ROWS = 10

def extract_xlsx(file_path: Path) -> ExtractedFile:
    wb = openpyxl.load_workbook(str(file_path), data_only=True, read_only=True)
    
    sheets_texts = []
    for sheet in wb.worksheets:
        sheets_texts.append(f"[Current Sheet: {sheet.title}]")
        row_count = 0
        for row in sheet.iter_rows(values_only=True):
            if row_count >= MAX_ROWS:
                break
            row_text = " | ".join(str(cell) for cell in row if cell is not None )
            if row_text:
                sheets_texts.append(row_text)
            row_count += 1
    
    full_text = "\n".join(sheets_texts).strip()
    
    
    return ExtractedFile(
        content=full_text if full_text else None,
        summary=None,
        metadata={
            "sheet_count": len(wb.worksheets),
            "sheet_names": wb.sheetnames          
        }
    )

# if __name__ == "__main__":
#     from pathlib import Path
#     result = extract_xlsx(Path("testing_dir1/sample_data.xlsx"))
#     print(result)